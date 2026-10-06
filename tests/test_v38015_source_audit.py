# v3.51.0 layout: tests/ -> put the repo root on sys.path, `import src` adds src/*
import sys, pathlib, json, tempfile, importlib.util, io, os, subprocess, contextlib, time
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
import src  # noqa: F401

"""v3.80.15 source audit (Excel day sheet vs Fubon source pages) - offline

The network goes through source_audit.http_get only; it is replaced by a fake
that serves synthetic Fubon pages (real page structure) and close lists.
  A. formulas mirrored from crawler.py / excel_report.py are still the crawler's
  B. Fubon page validation: identity, data date, truncated page
  C. row verdicts REAL / LOT_EST / AMT_EST / MISMATCH / UNVERIFIED
  D. sheet parser on a day sheet built by the REAL excel_report.build_day_sheet
     from rows merged by the REAL crawler_fetch.fetch_branch_combined
  E. end-to-end: ok / identity mismatch -> UNVERIFIED (not MISMATCH) /
     crawl-time page differences -> MISMATCH / time budget / closes down
  F. --email-line: ok / mismatch / incomplete / error / missing / stale json
  G. CLI always exits 0 and hands the status to the workflow
  H. daily-full.yml: step order, audit step settings, mail body, issue step
"""
import openpyxl
import yaml
from src.pipelines import crawler_fetch as cf
from src.audit import source_audit as sa
import excel_report as er

ROOT = pathlib.Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location('source_audit_cli', ROOT / 'scripts' / 'source_audit.py')
cli = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cli)
TMP = pathlib.Path(tempfile.mkdtemp())
TD = '20261005'
all_pass = True


def check(label, ok, detail=''):
    global all_pass
    print(f"  {'✅' if ok else '❌'} {label}" + (f"  ({detail})" if detail else ''))
    if not ok:
        all_pass = False


# ---------------------------------------------------------------- fixtures
def stock_row(code, name, v1, v2, gen=False):
    if gen:     # GenLink2stk form (most stocks on the real page)
        cell = f"<SCRIPT LANGUAGE=javascript>\n<!--\n\tGenLink2stk('AS{code}','{name}');\n//-->\n</SCRIPT>\n"
    else:       # Link2Stk form (ETFs and some stocks on the real page)
        cell = f"<a href=\"javascript:Link2Stk('{code}');\">{code}{name}</a>"
    return (f'<tr>\n<td class="t4t1" nowrap id="oAddCheckbox">\n{cell}\n</td>\n'
            f'<td class="t3n1" nowrap>{v1:,}</td>\n<td class="t3n1" nowrap>{v2:,}</td>\n'
            f'<td class="t3n1" nowrap>{v1 - v2:,}</td>\n</tr>\n')


def page_html(echo, date, buy, sell, truncated=False):
    head = ('<html><head><title>zgb0</title></head><body>\n<!-- ' + 'x' * 6000 + ' -->\n'
            f"<SCRIPT LANGUAGE=javascript><!--\n\tiID = '{echo}';\n"
            f"\tGenFundCorpCombo('{echo}','{echo}','frm');\n//--></SCRIPT>\n"
            f'<div class="t11">單位：仟元 ／ 資料日期：{date}</div>\n')
    body = ('<tr><td class="t2" colspan="4">買超</td></tr>\n'
            + ''.join(stock_row(*r, gen=(i % 2 == 1)) for i, r in enumerate(buy))
            + '<tr><td class="t2" colspan="4">賣超</td></tr>\n'
            + ''.join(stock_row(*r) for r in sell))
    if truncated:
        return head + body[:len(body) // 2]
    return head + body + '</table>\n</body>\n</html>\n'


# Source pages for TD (amounts in thousand TWD, lots). 9B25 and 9216 are each
# tracked by two masters; 9A9g / 9B25 are lettered (hex bno).
PAGES = {
    '9B25': {'B': {'buy': [('2454', '聯發科', 121327, 66364), ('3008', '大立光', 108245, 1155),
                           ('2330', '台積電', 50000, 0)],
                   'sell': [('2308', '台達電', 2145, 87995)]},
             'E': {'buy': [('00632R', '元大台灣50反1', 1900, 92), ('2303', '聯電', 315, 28),
                           ('2454', '聯發科', 24, 13), ('2330', '台積電', 34, 0)],
                   'sell': [('2308', '台達電', 3, 70)]}},
    '9666': {'B': {'buy': [('6488', '環球晶', 42108, 29851)], 'sell': []},
             'E': {'buy': [('6488', '環球晶', 43, 31)], 'sell': []}},
    '779W': {'B': {'buy': [('1303', '南亞', 35177, 24964), ('2409', '友達', 10034, 8423)], 'sell': []},
             'E': {'buy': [('1303', '南亞', 154, 109), ('2409', '友達', 283, 242)], 'sell': []}},
    '9216': {'B': {'buy': [('3105', '穩懋', 18450, 9747), ('6488', '環球晶', 23247, 10943)], 'sell': []},
             'E': {'buy': [('6488', '環球晶', 24, 11)], 'sell': []}},
    '9A9g': {'B': {'buy': [('2344', '華邦電', 18858, 4286)], 'sell': []},
             'E': {'buy': [], 'sell': []}},
}
CLOSES = {'2454': 1500.0, '3008': 6190.0, '2330': 1470.0, '2303': 156.25, '2308': 1000.0,
          '6488': 950.0, '1303': 227.5, '2409': 34.7, '2344': 172.5, '2603': 200.0}
OTC = {'3105': 290.5}                       # TPEx-only close
SENT = {cf.fubon_bno(b): b for b in PAGES}  # hex / plain bno -> branch code


def twse_doc(date=TD):
    return {'stat': 'OK', 'date': date, 'tables': [
        {'title': 'x', 'fields': ['證券代號', '證券名稱', '成交股數', '收盤價'],
         'data': [[c, 'n', '1,000', f'{p:,.2f}'] for c, p in CLOSES.items()] + [['9999', 'n', '0', '--']]}]}


def tpex_doc(date=TD):
    return {'date': date, 'stat': 'ok', 'tables': [
        {'title': 'x', 'fields': ['代號', '名稱', '收盤 ', '漲跌'],
         'data': [[c, 'n', f'{p:.2f}', '+0.1'] for c, p in OTC.items()]}]}


def qs(url):
    return dict(p.split('=', 1) for p in url.split('?', 1)[1].split('&'))


# ------------------------------------------------- crawler side (REAL merge)
def crawl(pages):
    """Run the REAL crawler_fetch.fetch_branch_combined on `pages` (d=1 URL),
    then crawler.py's close-based estimate block. -> branches_data."""
    class Resp:
        def __init__(self, html):
            self.status_code, self.content = 200, html.encode('big5')

    class Sess:
        def __init__(self):
            self.headers = {}

        def get(self, url, timeout=None):
            q = qs(url)
            assert q.get('d') == '1', url
            spec_ = pages[SENT[q['a']]][q['c']]
            return Resp(page_html(q['a'], TD, spec_['buy'], spec_['sell']))

    orig_sess, orig_sleep = cf.requests.Session, time.sleep
    cf.requests.Session, time.sleep = Sess, (lambda *_: None)
    try:
        out = []
        closes = {**OTC, **CLOSES}
        for bno in pages:
            d = cf.fetch_branch_combined(bno)
            assert d['error'] is None, d
            for s in d['buys'] + d['sells']:
                crawler_estimate(s, closes.get(s['code']))
            out.append({'code': bno, 'name': bno, 'date': d['date'],
                        'buys': d['buys'], 'sells': d['sells'], 'error': None})
        return out
    finally:
        cf.requests.Session, time.sleep = orig_sess, orig_sleep


def crawler_estimate(s, cp_close):
    """Same statements as crawler.py's v3.27.2/v3.80.1 block (section A pins them)."""
    if cp_close and cp_close > 0:
        buy_amt_k = s.get("buy_amt", 0) or 0
        sell_amt_k = s.get("sell_amt", 0) or 0
        buy_lot_raw = s.get("buy_lot", 0) or 0
        sell_lot_raw = s.get("sell_lot", 0) or 0
        estimated = False
        if buy_lot_raw == 0 and buy_amt_k > 0 and not s.get("lot_listed", False):
            s["buy_lot"] = max(1, round(buy_amt_k / cp_close)); estimated = True
        if sell_lot_raw == 0 and sell_amt_k > 0 and not s.get("lot_listed", False):
            s["sell_lot"] = max(1, round(sell_amt_k / cp_close)); estimated = True
        if buy_amt_k == 0 and buy_lot_raw > 0 and not s.get("amt_listed", False):
            s["buy_amt"] = int(round(buy_lot_raw * cp_close)); estimated = True
        if sell_amt_k == 0 and sell_lot_raw > 0 and not s.get("amt_listed", False):
            s["sell_amt"] = int(round(sell_lot_raw * cp_close)); estimated = True
        if estimated:
            s["net_amt"] = (s.get("buy_amt", 0) or 0) - (s.get("sell_amt", 0) or 0)
            s["net_lot"] = (s.get("buy_lot", 0) or 0) - (s.get("sell_lot", 0) or 0)


def build_xlsx(branches, name):
    """Day sheet from the REAL builder, saved to disk like the nightly file."""
    wb = openpyxl.Workbook()
    wb.active.title = '📋 今日 Dashboard'
    ws = wb.create_sheet(TD)
    # partial top-buyer stats -> the builder also writes its top warning row and
    # the footer row, as on a real night
    er.build_day_sheet(ws, branches, TD, precomputed_top_buyer={},
                       precomputed_stats={'attempted': 3, 'success': 1, 'fetch_fail': 2, 'fubon_success': 1})
    wb.create_sheet('20261002')            # older day sheet: must not be picked
    p = TMP / name
    wb.save(p)
    return p


# ------------------------------------------------- audit side (fake network)
CALLS = []


def make_get(pages, wrong_id=None, twse=None, tpex=None):
    def get(url, headers, timeout):
        CALLS.append(url)
        if 'twse.com.tw' in url:
            return twse or (200, json.dumps(twse_doc()).encode())
        if 'tpex.org.tw' in url:
            return tpex or (200, json.dumps(tpex_doc()).encode())
        q = qs(url)
        assert q['e'] == q['f'] == '2026-10-5', url
        bno = SENT[q['a']]
        spec_ = pages[bno][q['c']]
        echo = wrong_id if (wrong_id and bno == '9666') else q['a']
        return 200, page_html(echo, TD, spec_['buy'], spec_['sell']).encode('big5')
    return get


sa._sleep = lambda *_: None
quiet = lambda *a, **k: None


def audit(xlsx, get, **kw):
    CALLS.clear()
    return sa.run_audit(xlsx, get=get, log=quiet, **kw)


print("=" * 72)
print("  v3.80.15 來源比對 (Excel 日表 vs 富邦原始頁) — 離線")
print("=" * 72)

# ---------------------------------------------------------------- A
print("\nA. 估算公式與爬蟲 / 產表程式一致")
crawler_src = (ROOT / 'crawler.py').read_text(encoding='utf-8')
excel_src = (ROOT / 'src' / 'exports' / 'excel_report.py').read_text(encoding='utf-8')
for line in ['if cp_close and cp_close > 0:',
             'if buy_lot_raw == 0 and buy_amt_k > 0 and not s.get("lot_listed", False):',
             's["buy_lot"] = max(1, round(buy_amt_k / cp_close))',
             's["sell_lot"] = max(1, round(sell_amt_k / cp_close))',
             'if buy_amt_k == 0 and buy_lot_raw > 0 and not s.get("amt_listed", False):',
             's["buy_amt"] = int(round(buy_lot_raw * cp_close))',
             's["sell_amt"] = int(round(sell_lot_raw * cp_close))']:
    check(f"crawler.py 仍是: {line}", line in crawler_src)
check("excel_report.py 仍是: round(buy_amt_k / 10) if buy_amt_k else 0",
      'buy_amt_w = round(buy_amt_k / 10) if buy_amt_k else 0' in excel_src
      and 'sell_amt_w = round(sell_amt_k / 10) if sell_amt_k else 0' in excel_src)
check("crawler_fetch 張數/金額只在同一區 (買超) 整頁查表",
      'amt_map = {r["code"]: r for r in amt_rows}' in (ROOT / 'src/pipelines/crawler_fetch.py').read_text(encoding='utf-8'))
check("估算張數: 1155 仟元 / 6190 → 1 張 (max 1), 0 → 0", sa.est_lots(1155, 6190.0) == 1 and sa.est_lots(0, 6190.0) == 0)
check("估算金額: 28 張 x 156.25 = 4375 仟元 → 438 萬 (同 Python round)",
      sa.est_amt_k(28, 156.25) == 4375 and sa.to_wan(4375) == 438 and sa.to_wan(0) == 0)

# ---------------------------------------------------------------- B
print("\nB. 富邦頁面驗證 (身分 / 資料日期 / 截斷)")
hexid = cf.fubon_bno('9B25')
good = page_html(hexid, TD, PAGES['9B25']['B']['buy'], PAGES['9B25']['B']['sell'])
p = sa.parse_fubon_page(good, hexid, TD)
check("正常頁: 買超 3 檔 / 賣超 1 檔, 兩種連結寫法都讀到",
      [r['code'] for r in p['buy']] == ['2454', '3008', '2330'] and [r['code'] for r in p['sell']] == ['2308'])


def rejects(html, sent, word):
    try:
        sa.parse_fubon_page(html, sent, TD)
    except sa.SourceError as e:
        return word in str(e)
    return False


check("頁面自報別的分點 → identity mismatch",
      rejects(page_html('9B25', TD, [('2330', 'x', 1, 0)], []), hexid, 'identity mismatch'))
check("頁面沒有分點 id 標記 → 拒收",
      rejects(good.replace('frm', 'xxx'), hexid, 'no branch-id marker'))
check("資料日期不是交易日 → date mismatch",
      rejects(page_html(hexid, '20261002', [('2330', 'x', 1, 0)], []), hexid, 'date mismatch'))
check("頁面被截斷 (沒有 </html>) → 拒收",
      rejects(page_html(hexid, TD, PAGES['9B25']['B']['buy'], [], truncated=True), hexid, 'truncated'))
check("頁面過小 → 拒收", rejects('<html></html>', hexid, 'too small'))
check("URL: hex 代號 + 指定單日 e=f=2026-10-5",
      sa.fubon_url('9A9g', 'E', TD).endswith('a=0039004100390067&b=0039004100390067&c=E&e=2026-10-5&f=2026-10-5'))

# ---------------------------------------------------------------- C
print("\nC. 逐列判定")
X = lambda bl, sl, bw, sw: {'buy_lot': bl, 'sell_lot': sl, 'buy_wan': bw, 'sell_wan': sw}
v = lambda *a: sa.classify_row(*a)[0]
check("REAL: 兩頁都有, 張數與金額都等於頁面", v(X(24, 13, 12133, 6636), (121327, 66364), (24, 13), 1500.0) == 'REAL')
check("REAL 但張數不同 → MISMATCH", v(X(25, 13, 12133, 6636), (121327, 66364), (24, 13), 1500.0) == 'MISMATCH')
check("LOT_EST: 張數頁沒有, 張數 = 金額/收盤 (17, max(1,0)=1)",
      v(X(17, 1, 10824, 116), (108245, 1155), None, 6190.0) == 'LOT_EST')
check("LOT_EST 估算錯 → MISMATCH", v(X(18, 1, 10824, 116), (108245, 1155), None, 6190.0) == 'MISMATCH')
check("AMT_EST: 金額頁沒有, 金額 = 張數 x 收盤", v(X(315, 28, 4922, 438), None, (315, 28), 156.25) == 'AMT_EST')
check("AMT_EST 估算錯 → MISMATCH", v(X(315, 28, 4921, 438), None, (315, 28), 156.25) == 'MISMATCH')
check("兩頁都沒有這檔 → MISMATCH", sa.classify_row(X(1, 0, 10, 0), None, None, 10.0)[:3:2] == ('MISMATCH', sa.R_NOT_ON_PAGE))
check("估算列沒有官方收盤價 → UNVERIFIED (不是 MISMATCH)", v(X(17, 1, 10824, 116), (108245, 1155), None, None) == 'UNVERIFIED')
r = sa.classify_row(X(289, 243, 1003, 842), (10034, 8423), (283, 242), 34.7)
check("張數被估算覆蓋 (頁上其實有) → MISMATCH 並註明是估算值", r[0] == 'MISMATCH' and sa.H_LOTS_EST in r[2], r[2])

# ---------------------------------------------------------------- D
print("\nD. 解析真正 build_day_sheet 產出的日表 (REAL crawler merge → REAL builder → 存檔再讀)")
branches = crawl(PAGES)
xlsx_ok = build_xlsx(branches, 'ok.xlsx')
wb = openpyxl.load_workbook(xlsx_ok, read_only=True)
check("最新日表 = 20261005 (忽略 Dashboard 與較舊日表)", sa.newest_day_sheet(wb.sheetnames) == TD, wb.sheetnames)
rows = sa.parse_day_sheet(wb[TD])
wb.close()
by_code = {b['code']: b for b in branches}
want = []
for m in er.MASTER_MAPPING:
    for bno, bname in m['branches']:
        for s in er._top_stocks_for_branch(by_code.get(bno, {}), sniper_mode=er._is_sniper_master(m['name']),
                                           n_top=er._branch_stocks_size(bno)):
            want.append((m['name'], bname, bno, s['code'], s['buy_lot'], s['sell_lot'],
                         round(s['buy_amt'] / 10) if s['buy_amt'] else 0,
                         round(s['sell_amt'] / 10) if s['sell_amt'] else 0))
got = [(x['master'], x['branch'], x['bno'], x['code'], x['buy_lot'], x['sell_lot'], x['buy_wan'], x['sell_wan'])
       for x in rows]
check("每一列的 高手/分點/代號/股票/E~H 都讀對 (16 列)", got == want and len(got) == 16, f"{len(got)} vs {len(want)}")
check("同一分點掛兩位高手 (9B25 民哥+強森, 9216 林滄海+陳族元) 各自成列",
      {x['master'] for x in rows if x['bno'] == '9B25'} == {'民哥', '強森'}
      and {x['master'] for x in rows if x['bno'] == '9216'} == {'林滄海', '陳族元'})
check("ETF (00632R) 與淨賣股 (2308) 不在日表", not ({'00632R', '2308'} & {x['code'] for x in rows}))
check("'常下分點' 標頭 (C 欄 '分點代號') 不被當成分點代號", all(x['bno'] in PAGES for x in rows))
wb = openpyxl.load_workbook(xlsx_ok)
col_a = [wb[TD].cell(r, 1).value or '' for r in range(1, wb[TD].max_row + 1)]
check("日表含頂端警示列 (⚠️) 與 ⓘ 尾列 (真實夜間的樣子), 解析時略過",
      col_a[0].startswith('⚠') and col_a[-1].startswith('ⓘ') and all(x['master'] in er.TRACKED_MASTERS for x in rows))

closes_d, _ = sa.fetch_closes(TD, make_get(PAGES))
res_d, _ = sa.audit_rows(rows, TD, closes_d, get=make_get(PAGES), log=quiet)
per = {(x['bno'], x['code']): r['verdict'] for x, r in zip(rows, res_d)}
WANT = {('9B25', '2454'): 'REAL', ('9B25', '3008'): 'LOT_EST', ('9B25', '2330'): 'REAL',
        ('9B25', '2303'): 'AMT_EST', ('9666', '6488'): 'REAL', ('779W', '1303'): 'REAL',
        ('779W', '2409'): 'REAL', ('9216', '6488'): 'REAL', ('9216', '3105'): 'LOT_EST',
        ('9A9g', '2344'): 'LOT_EST'}
check("每列判定正確 (3105 上櫃股用 TPEx 收盤估張數, 2303 只在張數頁 → 估金額)", per == WANT, per)

# ---------------------------------------------------------------- E
print("\nE. 端到端 (假網路)")
d1 = audit(xlsx_ok, make_get(PAGES))
vd = d1['verdicts']
check("全部相符 → status ok, 真實 9 / 估算張數 5 / 估算金額 2",
      d1['status'] == 'ok' and (vd['REAL'], vd['LOT_EST'], vd['AMT_EST'], vd['MISMATCH'], vd['UNVERIFIED']) == (9, 5, 2, 0, 0),
      (d1['status'], vd, d1.get('error')))
fub = [u for u in CALLS if 'fubon' in u]
check("每個分點只抓一次 (5 分點 x 2 頁 = 10 次, 重複掛名不重抓)", len(fub) == 10, len(fub))
check("上櫃股收盤 (TPEx) 也用上: 3105 為 LOT_EST", d1['closes_loaded'] == len(CLOSES) + 1)

d2 = audit(xlsx_ok, make_get(PAGES, wrong_id='9667'))
check("9666 原始頁自報別的分點 → 該分點 UNVERIFIED, 不算 MISMATCH",
      d2['status'] == 'incomplete' and d2['verdicts']['MISMATCH'] == 0 and d2['verdicts']['UNVERIFIED'] == 1
      and [b['bno'] for b in d2['unverified_branches']] == ['9666']
      and 'identity mismatch' in d2['unverified_branches'][0]['reason'], (d2['status'], d2['verdicts'], d2['unverified_branches']))
check("身分不符會重試一次 (9666 金額頁 2 次)", sum(1 for u in CALLS if 'a=9666&' in u) == 2)

crawl_bad = json.loads(json.dumps(PAGES))
crawl_bad['779W']['E']['buy'] = [crawl_bad['779W']['E']['buy'][0]]          # crawl-time lots page cut short
crawl_bad['9666']['B']['buy'].append(['2603', '長榮', 30000, 0])             # crawl-time page had another stock
crawl_bad = {b: {m: {r: [tuple(x) for x in rows_] for r, rows_ in md.items()} for m, md in bd.items()}
             for b, bd in crawl_bad.items()}
xlsx_bad = build_xlsx(crawl(crawl_bad), 'bad.xlsx')
d3 = audit(xlsx_bad, make_get(PAGES))
mm = {(m['bno'], m['code']): m for m in d3['mismatches']}
check("Excel 與指定日期原始頁不同 → status mismatch, 2 列", d3['status'] == 'mismatch' and d3['mismatch_total'] == 2,
      (d3['status'], d3['verdicts']))
m = mm.get(('779W', '2409'), {})
check("友達: Excel 289/243 張 (估算) vs 頁面 283/242 → 列出 Excel 值、應為值、來源值",
      m.get('excel', {}).get('buy_lot') == 289 and m.get('expected', {}).get('buy_lot') == 283
      and m.get('source', {}).get('lots') == [283, 242] and sa.H_LOTS_EST in m.get('reason', ''), m)
check("長榮: 原始頁沒有這檔 → MISMATCH", mm.get(('9666', '2603'), {}).get('reason') == sa.R_NOT_ON_PAGE)

ticks = {'t': 0.0}
def clock():
    return ticks['t']
def slow_get(url, headers, timeout):
    ticks['t'] += 50
    return make_get(PAGES)(url, headers, timeout)
d4 = audit(xlsx_ok, slow_get, clock=clock, sleep=lambda *_: None, budget_s=220)
check("時間預算用完 → 其餘分點 UNVERIFIED (time budget), 已抓的照常判定",
      d4['status'] == 'incomplete' and d4['verdicts']['REAL'] + d4['verdicts']['LOT_EST'] + d4['verdicts']['AMT_EST'] == 8
      and [b['bno'] for b in d4['unverified_branches']] == ['9666', '779W', '9216', '9A9g']
      and all('budget' in b['reason'] for b in d4['unverified_branches']), (d4['verdicts'], d4['unverified_branches']))

d5 = audit(xlsx_ok, make_get(PAGES, twse=(500, b''), tpex=(200, json.dumps(tpex_doc('20261002')).encode())))
check("收盤價抓不到 → 估算列 UNVERIFIED (7 列), 真實列照常相符, 不是 MISMATCH",
      d5['status'] == 'incomplete' and d5['verdicts']['REAL'] == 9 and d5['verdicts']['UNVERIFIED'] == 7
      and d5['unverified_rows_close'] == 7 and len(d5['close_errors']) == 2, (d5['verdicts'], d5['close_errors']))

d6 = sa.run_audit(TMP / 'missing.xlsx', get=make_get(PAGES), log=quiet)
check("Excel 不存在 → status error (不丟例外)", d6['status'] == 'error' and d6['error'], d6['error'])

# ---------------------------------------------------------------- F
print("\nF. Email 第一行 (--email-line)")
idx = TMP / 'index.json'
idx.write_text(json.dumps({'latest': TD}), encoding='utf-8')


def line_for(doc):
    p = TMP / 'audit.json'
    if doc is None:
        p.unlink(missing_ok=True)
    else:
        p.write_text(json.dumps(doc, ensure_ascii=False), encoding='utf-8')
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = cli.main(['--email-line', '--out', str(p), '--index', str(idx)])
    return rc, buf.getvalue().strip()


cases = [
    (d1, '✅ 來源比對 10/05：16 列全部相符（真實 9、估算張數 5、估算金額 2）'),
    (d3, '🚨 來源比對 10/05：2 列不符，詳見 GitHub Issue'),
    (d2, '⚠️ 來源比對 10/05：1 個分點抓不到原始頁，未比對 1 列'),
    (d5, '⚠️ 來源比對 10/05：7 列缺官方收盤價，未比對 7 列'),
    ({**d1, 'status': 'error', 'error': 'x'}, '⚠️ 來源比對 10/05：比對程式出錯，詳見 Actions 紀錄'),
    (None, '⚠️ 來源比對：今天沒有跑完'),
    ({**d1, 'trade_date': '20261002'}, '⚠️ 來源比對：今天沒有跑完'),
]
for doc, want_line in cases:
    rc, line = line_for(doc)
    check(f"{want_line}", rc == 0 and line == want_line, line)

# ---------------------------------------------------------------- G
print("\nG. CLI 一律 exit 0, status 交給 workflow")
out_json, gh_out = TMP / 'cli.json', TMP / 'gh_output.txt'
os.environ['GITHUB_OUTPUT'] = str(gh_out)
sa.http_get = make_get(PAGES)
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    rc = cli.main(['--xlsx', str(xlsx_bad), '--out', str(out_json)])
doc = json.loads(out_json.read_text(encoding='utf-8'))
check("不符 → exit 0, json status mismatch, ::error 註記, GITHUB_OUTPUT status=mismatch",
      rc == 0 and doc['status'] == 'mismatch' and '::error title=來源比對不符::' in buf.getvalue()
      and 'status=mismatch' in gh_out.read_text(encoding='utf-8'), buf.getvalue()[-200:])
check("json 欄位齊全", all(k in doc for k in ('trade_date', 'checked_at', 'status', 'rows', 'verdicts',
                                             'mismatches', 'unverified_branches', 'error')))
gh_out.unlink()
env = {**os.environ, 'PYTHONIOENCODING': 'utf-8', 'GITHUB_OUTPUT': str(gh_out)}
pr = subprocess.run([sys.executable, str(ROOT / 'scripts' / 'source_audit.py'), '--xlsx', str(TMP / 'nope.xlsx'),
                     '--out', str(TMP / 'err.json')], capture_output=True, text=True, encoding='utf-8', env=env)
err_doc = json.loads((TMP / 'err.json').read_text(encoding='utf-8'))
check("Excel 不存在 (子程序) → exit 0, status error, ::warning, GITHUB_OUTPUT status=error",
      pr.returncode == 0 and err_doc['status'] == 'error' and '::warning' in pr.stdout
      and 'status=error' in gh_out.read_text(encoding='utf-8'), (pr.returncode, pr.stdout[-200:], pr.stderr[-200:]))
pr = subprocess.run([sys.executable, str(ROOT / 'scripts' / 'source_audit.py'), '--no-such-flag'],
                    capture_output=True, text=True, encoding='utf-8', env=env)
check("參數錯 → 仍 exit 0", pr.returncode == 0, pr.returncode)
del os.environ['GITHUB_OUTPUT']

# ---------------------------------------------------------------- H
print("\nH. daily-full.yml")
wf = yaml.safe_load(open(ROOT / '.github/workflows/daily-full.yml', encoding='utf-8'))
job = wf['jobs']['crawl']
steps = job['steps']
names = [s.get('name') for s in steps]
OLD = ['Checkout repo', 'Setup Python', 'Install dependencies', 'Download previous DB', 'Run full crawler',
       'Refresh attstock disposal', 'Update Phase 3.2 rolling backtest + regen Excel', 'Upload DB artifact',
       'Commit and push data', 'Extract mobile summary', 'Send daily summary email',
       'Email failure alert (GitHub Issue)']
check("既有步驟名稱與順序不變", [n for n in names if n in OLD] == OLD, names)
AUD = 'Source audit (Excel vs Fubon pages)'
ia = names.index(AUD) if AUD in names else -1
check("來源比對緊接在最後一次產表之後、上傳 DB 之前",
      ia > 0 and names[ia - 1] == 'Update Phase 3.2 rolling backtest + regen Excel' and names[ia + 1] == 'Upload DB artifact')
st = steps[ia] if ia >= 0 else {}
check("continue-on-error / timeout 10 分 / id audit / 跑 scripts/source_audit.py",
      st.get('continue-on-error') is True and st.get('timeout-minutes') == 10 and st.get('id') == 'audit'
      and 'python scripts/source_audit.py' in st.get('run', ''), st)
run = steps[names.index('Extract mobile summary')]['run']
check("Email 內文第一行 = --email-line, 接一個空行, 再接手機摘要",
      'scripts/source_audit.py --email-line' in run
      and run.index('echo "summary<<MOBILE_EOF"') < run.index('echo "$AUDIT_LINE"')
      < run.index('echo ""') < run.index('echo "$SUMMARY"'))
ISS = 'Source audit mismatch alert (GitHub Issue)'
ii = names.index(ISS) if ISS in names else -1
iss = steps[ii] if ii >= 0 else {}
cond = str(iss.get('if', ''))
check("開 Issue 步驟在寄信之後", ii > names.index('Send daily summary email'))
check("條件: status == mismatch 且 data_changed == true",
      "steps.audit.outputs.status == 'mismatch'" in cond and "steps.commit.outputs.data_changed == 'true'" in cond, cond)
script = (iss.get('with') or {}).get('script', '')
check("actions/github-script, label source-audit, 同標題 open issue 存在就跳過",
      str(iss.get('uses', '')).startswith('actions/github-script') and "labels: ['source-audit']" in script
      and "state: 'open'" in script and 'i.title === title' in script and 'mismatches' in script)
check("crawl job 有 issues: write (開 Issue 需要)", (job.get('permissions') or {}).get('issues') == 'write'
      and (job.get('permissions') or {}).get('contents') == 'write')
check("data/audit/ 不進 git (每次都會改寫, 不可讓 data_changed 永遠為真)",
      'data/audit/' in (ROOT / '.gitignore').read_text(encoding='utf-8'))

print("\nI. v3.80.18 櫃買網站抓不到 → 改用櫃買 OpenAPI (2026-10-06 雲端: 34 列缺收盤價)")
def openapi_doc(roc='1151005'):
    return [{'Date': roc, 'SecuritiesCompanyCode': c, 'Close': f'{p:.2f}'} for c, p in OTC.items()]
def close_get(openapi):
    def get(url, headers, timeout):
        if 'twse.com.tw' in url:
            return 200, json.dumps(twse_doc()).encode()
        if 'openapi' in url:
            return openapi
        return 500, b''                          # TPEx website down
    return get
c_ok, e_ok = sa.fetch_closes(TD, close_get((200, ('﻿' + json.dumps(openapi_doc())).encode('utf-8'))))
check("網站 500 → OpenAPI (含 BOM) 補上櫃收盤價, 不算錯誤", c_ok.get('3105') == 290.5 and e_ok == [], (c_ok.get('3105'), e_ok))
c_bad, e_bad = sa.fetch_closes(TD, close_get((200, json.dumps(openapi_doc('1151002')).encode())))
check("OpenAPI 是別天的 → 不採用, 錯誤寫明兩個來源", '3105' not in c_bad and len(e_bad) == 1
      and 'OpenAPI fallback' in e_bad[0] and 'date mismatch' in e_bad[0], e_bad)
check("民國日期換算 20261005 → 1151005", sa._roc_date('20261005') == '1151005')
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    cli.annotate({'trade_date': TD, 'status': 'incomplete', 'verdicts': {'UNVERIFIED': 34},
                  'unverified_branches': [], 'unverified_rows_close': 34,
                  'close_errors': ['TPEx: HTTP 500; OpenAPI fallback: HTTP 500']})
ann = buf.getvalue()
check("雲端警告寫出真正原因 (缺官方收盤價 + 來源錯誤), 不再寫「0 個分點抓不到」",
      '34 列缺官方收盤價' in ann and 'HTTP 500' in ann and '0 個分點' not in ann, ann.strip())

print()
print("─" * 72)
print(f"  整體: {'✅ ALL PASS' if all_pass else '❌ HAS FAIL'}")
sys.exit(0 if all_pass else 1)
