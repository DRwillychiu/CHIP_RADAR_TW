# v3.51.0 layout: tests/ -> put the repo root on sys.path, `import src` adds src/*
import sys, pathlib, json, tempfile
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
import src  # noqa: F401

"""v3.80.21 焦家 master + 🔁 焦家匯撥 transfer-watch sheet (offline, synthetic data)

  A. registry: 585b 統一-內湖 (master 焦家), 984K / 989N co_masters, style unknown,
     MASTER_MAPPING block + color, mapping validation 0 warnings, hex bno vs 585B
  B. consensus counts the PRIMARY master only: 984K buying does not add 焦家
  C. sheet from fake today + fake stored days + fake page fetch: labels, values,
     dash, approx sign, N-day total, footnote for missing data
  D. cache: second build (fetch allowed) and the final regen (fetch=False) never
     refetch; request cap
  E. stored day files go through quarantine.filter_day; a quarantined /
     failed branch-day is read from the page instead
  F. tab order Dashboard -> Pinned -> transfer sheet -> date sheets
  G. real fetch path (source_audit) with a fake http_get: hex id, identity check
"""
import openpyxl
import excel_report as er
from src.core import branches as reg
from src.pipelines import crawler_fetch as cf
from src.pipelines.crawler_output import encrypt_data

tw = er._transfer_watch()
sa = tw.sa
all_pass = True


def check(label, ok, detail=''):
    global all_pass
    print(f"  {'✅' if ok else '❌'} {label}" + (f"  ({detail})" if detail else ''))
    if not ok:
        all_pass = False


print("=" * 72)
print("  v3.80.21 焦家 + 匯撥追蹤 (offline)")
print("=" * 72)

# ── A ──
print("\nA. 名冊")
b585 = reg.get_branch_by_code('585b')
check("585b = 統一-內湖, master 焦家", b585 and b585['name'] == '統一-內湖' and b585['master'] == '焦家', b585)
check("984K master 強森, co_masters 巨人傑 + 焦家",
      reg.get_branch_by_code('984K')['master'] == '強森'
      and reg.get_branch_by_code('984K')['co_masters'] == ['巨人傑', '焦家'])
check("989N master 強森, co_masters 焦家",
      reg.get_branch_by_code('989N')['master'] == '強森'
      and reg.get_branch_by_code('989N')['co_masters'] == ['焦家'])
check("風格 = unknown (使用者尚未指定)", reg.MASTER_STYLES.get('焦家') == ['unknown'])
blk = next((m for m in er.MASTER_MAPPING if m['name'] == '焦家'), None)
check("MASTER_MAPPING 焦家區塊 = 984K / 989N / 585b",
      blk and blk['branches'] == [('984K', '元大-館前'), ('989N', '元大-內湖'), ('585b', '統一-內湖')], blk)
check("MASTER_BLOCK_COLORS 有 焦家", '焦家' in er.MASTER_BLOCK_COLORS)
w = er._validate_master_mapping_vs_branches()
check("MASTER_MAPPING vs branches.py 0 warnings", w == [], w)
check("585b 以 hex 送出, 與大小寫雙胞胎 585B (統一-永和) 不同",
      cf.fubon_bno('585b') == '0035003800350062' and cf.fubon_bno('585B') == '0035003800350042')
check("585b 只出現一次", sum(1 for b in reg.WATCHED_BRANCHES if b['code'] == '585b') == 1)

# ── B ──
print("\nB. 共識只算主要 master")


def rec(code, buys):
    b = dict(reg.get_branch_by_code(code))
    return {**b, 'buys': buys, 'sells': [], 'error': None}


def buy(code, net_k):
    return {'code': code, 'name': code, 'buy_amt': net_k + 10, 'sell_amt': 10}


others = [m['name'] for m in er.MASTER_MAPPING if m['name'] not in ('焦家', '強森')][:9]
fake = [{'code': f'X{i:02d}', 'name': f'X{i:02d}', 'master': m, 'buys': [buy('2492', 100)], 'sells': []}
        for i, m in enumerate(others)]
picks = {p['code']: p for p in er._compute_consensus_count(fake + [rec('984K', [buy('2492', 500)])])}
m2492 = picks.get('2492', {}).get('masters', set())
check("984K 買 2492 → master 集合含 強森, 不含 焦家 (co_master 不算)",
      '強森' in m2492 and '焦家' not in m2492, sorted(m2492))
picks = {p['code']: p for p in er._compute_consensus_count(
    fake + [rec('984K', [buy('2492', 500)]), rec('989N', [buy('2492', 50)])])}
check("再加 989N 也不會加入 焦家", '焦家' not in picks['2492']['masters'])
picks = {p['code']: p for p in er._compute_consensus_count(fake + [rec('585b', [buy('2492', 50)])])}
check("585b 買 2492 → 焦家 才算進去", '焦家' in picks['2492']['masters'], sorted(picks['2492']['masters']))

# ── C ── fake world
print("\nC. 匯撥追蹤頁 (假資料)")
TD = '20261005'
PAST = ['20261002', '20261001', '20260930', '20260929']
G = tw.TRANSFER_WATCH_GROUPS[0]


def stock(code, bl, sl, ba, sm, **kw):
    return {'code': code, 'name': code, 'buy_lot': bl, 'sell_lot': sl, 'buy_amt': ba, 'sell_amt': sm,
            'lot_listed': True, 'amt_listed': True, 'lot_source': 'twse', **kw}


def filler(n, prefix):    # pad a list to n rows (crawler keeps TOP_N rows per page region)
    return [stock(f'{prefix}{i:03d}', 1, 0, 10, 0) for i in range(n)]


def branch(code, buys, sells, date, **kw):
    return {**reg.get_branch_by_code(code), 'buys': buys, 'sells': sells, 'date': date, 'error': None, **kw}


# today: 984K full lists (2492 real +28, 4919 absent -> page), 989N short lists
# (complete: 6173 absent = known dash), 585b missing from the crawl -> page
today = [
    branch('984K', [stock('2492', 95, 67, 35550, 25591)] + filler(40, 'A'), filler(35, 'S'), TD),
    branch('989N', [stock('4919', 20, 3, 3082, 487), stock('2492', 0, 0, 35230, 32675,
                    lot_listed=False, lot_source='estimated_from_close')], [], TD),
]
today[1]['buys'][1].update(buy_lot=90, sell_lot=84)    # crawler estimate stored (approx)
STORED = {
    '20261002': [branch('984K', [stock('4919', 22, 73, 3546, 11110)] + filler(30, 'A'), [], '20261002'),
                 branch('989N', [], [], '20261002'),                       # no trades that day
                 branch('585b', [stock('2492', 18, 38, 6314, 13642)], [], '20261002')],
    '20261001': [branch('984K', [], [], '20261001', error='HTTP 500'),     # failed branch -> page
                 branch('989N', [stock('6173', 5, 1, 1500, 300)], [], '20261001'),
                 branch('585b', [], [], '20261001', quarantined=True, error='quarantined: x')],
    '20260930': [branch('984K', [], [], '20260930'), branch('989N', [], [], '20260930'),
                 branch('585b', [], [], '20260930')],
    '20260929': [branch('984K', [], [], '20260929'), branch('989N', [], [], '20260929'),
                 branch('585b', [], [], '20260929')],
}
# date-pinned pages: {(bno, date): {code: (amt, lots)}}
PAGES = {
    ('984K', TD): {'4919': (None, None), '6173': ((900, 0), (3, 0))},
    ('585b', TD): {'2492': ((4988, 7682), (13, 20)), '4919': ((6240, 151), (41, 1)), '6173': (None, None)},
    ('984K', '20261001'): {'2492': ((11809, 22074), (38, 70)), '4919': ((11657, 5426), None)},
    ('585b', '20261001'): {'6173': ((6632, 9383), (22, 32))},
}
CLOSES = {TD: {'2492': 368.5, '4919': 155.5, '6173': 292.0}, '20261001': {'4919': 152.0}}
CHANGE = {TD: {'2492': 1.94, '4919': -4.6, '6173': 0.0}}


def make_dir():
    d = pathlib.Path(tempfile.mkdtemp())
    (d / 'archive').mkdir()
    for x in PAST[:-1]:
        (d / f'{x}.json').write_text('{}', encoding='utf-8')
    (d / 'archive' / f'{PAST[-1]}.json').write_text('{}', encoding='utf-8')
    (d / '20260925.json').write_text('{}', encoding='utf-8')     # 6th day: outside the window
    sh = {'stocks': {c: {'daily': {day: {'close': CLOSES[day][c], 'change_pct': CHANGE.get(day, {}).get(c)}
                                   for day in CLOSES if c in CLOSES[day]}}
                     for c in ('2492', '4919', '6173')}}
    (d / 'stock_history.json').write_text(json.dumps(sh), encoding='utf-8')
    return d


calls = []


def fake_fetch(bno, date, pacer):
    calls.append((bno, date))
    spec = PAGES.get((bno, date))
    if spec is None:
        raise sa.SourceError('page has no data date')
    B = {'buy': {}, 'sell': {}}
    E = {'buy': {}, 'sell': {}}
    for code, (amt, lots) in spec.items():
        if amt:
            B['buy' if amt[0] >= amt[1] else 'sell'][code] = amt
        if lots:
            E['buy' if (amt or lots)[0] >= (amt or lots)[1] else 'sell'][code] = lots
    return {'B': B, 'E': E}


reads = []


def fake_read(data_dir, date, password):
    reads.append(date)
    return {'trade_date': date, 'branches': STORED.get(date, [])}


ORIG_FETCH, ORIG_READ = tw.fetch_branch_day, tw.read_day
tw.fetch_branch_day, tw.read_day = fake_fetch, fake_read
DIR = make_dir()
wb = openpyxl.Workbook()
names = er._build_transfer_sheets(wb, today, TD, DIR, fetch=True)
check("建出 1 張: 🔁 焦家匯撥", names == ['🔁 焦家匯撥'], names)
check("視窗 = 今天 + 4 個已存日 (含 archive, 不含第 6 日)", sorted(reads) == sorted(PAST), reads)
ws = wb['🔁 焦家匯撥']


def find_row(col, value, start=1):
    return next((r for r in range(start, ws.max_row + 1) if ws.cell(r, col).value == value), None)


def val(r, c):
    x = ws.cell(r, c)
    return (None if x.value == '—' else x.value), ('≈' in (x.number_format or ''))


check("標題 / 日期", ws['B2'].value == '焦家匯撥追蹤' and str(ws['H2'].value).startswith('2026/10/05'),
      (ws['B2'].value, ws['H2'].value))
check("副標 = 館前買進 → 匯撥至 內湖", ws['B3'].value.startswith('元大-館前 買進 → 匯撥至 元大-內湖、統一-內湖'),
      ws['B3'].value)
H = find_row(2, '個股')
hdr = {ws.cell(H, c).value: c for c in range(2, 12) if ws.cell(H, c).value}
check("今日表頭 (以標籤定位)", list(hdr) == ['個股', '代號', '元大-館前', '元大-內湖', '統一-內湖', '三戶合計', '收盤', '漲跌%'],
      list(hdr))
rows = {str(ws.cell(r, hdr['代號']).value): r for r in range(H + 1, H + 4)}
r = rows.get('2492')
check("2492 館前 = 95-67 = +28 (真實)", val(r, hdr['元大-館前']) == (28, False), val(r, hdr['元大-館前']))
check("2492 內湖 = 爬蟲估算 90-84 = ≈6", val(r, hdr['元大-內湖']) == (6, True), val(r, hdr['元大-內湖']))
check("2492 統一-內湖 (爬蟲沒有 → 指定日期頁) 13-20 = -7", val(r, hdr['統一-內湖']) == (-7, False))
check("2492 三戶合計 = 27, 含估算 → ≈", val(r, hdr['三戶合計']) == (27, True), val(r, hdr['三戶合計']))
check("2492 收盤 368.5 / 漲跌 +1.94%", ws.cell(r, hdr['收盤']).value == 368.5
      and abs(ws.cell(r, hdr['漲跌%']).value - 0.0194) < 1e-9)
r = rows.get('4919')
check("4919 館前: 不在前 30 名 → 查頁 → 頁上也沒有 → —", ws.cell(r, hdr['元大-館前']).value == '—')
r = rows.get('6173')
check("6173 館前: 頁上只有金額頁… 有張數 3 → +3", val(r, hdr['元大-館前']) == (3, False), val(r, hdr['元大-館前']))
check("6173 內湖: 名單完整 (<30 列) 且沒有 → — (不必查頁)", ws.cell(r, hdr['元大-內湖']).value == '—'
      and ('989N', TD) not in calls)
check("紅 = 正 / 綠 = 負", ws.cell(rows['2492'], hdr['元大-館前']).font.color.rgb == 'FFC62828'
      and ws.cell(rows['2492'], hdr['統一-內湖']).font.color.rgb == 'FF2E7D32')
check("KPI 卡: 華新科 今日三戶合計 ≈+27 張", ws['B6'].value == '≈+27 張' and ws['B5'].value == '華新科 今日三戶合計',
      (ws['B5'].value, ws['B6'].value))

# 5-day table of 4919
r0 = find_row(2, '新唐（4919）')
h5 = {ws.cell(r0 + 1, c).value: c for c in range(2, 7)}
drows = {ws.cell(rr, 2).value[:5]: rr for rr in range(r0 + 2, r0 + 7)}
check("近 5 日日期 = 10/05 … 09/29 (新到舊)", list(drows) == ['10/05', '10/02', '10/01', '09/30', '09/29'],
      list(drows))
check("10/02 館前 22-73 = -51 (已存日)", val(drows['10/02'], h5['元大-館前']) == (-51, False))
check("10/02 內湖: 當天沒有進出 → —", ws.cell(drows['10/02'], h5['元大-內湖']).value == '—')
check("10/01 館前: 存檔失敗分點 → 查頁 → 張數頁沒有 → 金額÷收盤 ≈ 77-36 = ≈41",
      val(drows['10/01'], h5['元大-館前']) == (41, True), val(drows['10/01'], h5['元大-館前']))
tot_r = find_row(2, '5 日合計', r0)
check("5 日合計 館前 = -51 + 41 = ≈-10", val(tot_r, h5['元大-館前']) == (-10, True), val(tot_r, h5['元大-館前']))
check("5 日合計列粗體 + 上框線", ws.cell(tot_r, h5['元大-館前']).font.b
      and ws.cell(tot_r, 2).border.top.style == 'medium')
r6 = find_row(2, '信昌電（6173）')
d6 = {ws.cell(rr, 2).value[:5]: rr for rr in range(r6 + 2, r6 + 7)}
check("10/01 統一-內湖: 隔離中的分點日 → 查頁 22-32 = -10", val(d6['10/01'], h5['統一-內湖']) == (-10, False))
check("10/02 統一-內湖: 存檔有該分點但不在名單 → —", ws.cell(d6['10/02'], h5['統一-內湖']).value == '—'
      and ('585b', '20261002') not in calls)
notes = [ws.cell(rr, 2).value for rr in range(1, ws.max_row + 1)
         if isinstance(ws.cell(rr, 2).value, str) and ws.cell(rr, 2).value.startswith(('匯撥', '三個', '資料'))]
check("註腳: 匯撥不經市場 / 三個分點有其他客戶 / 資料來源",
      notes[0].startswith('匯撥是帳戶之間移轉、不經市場') and notes[1].startswith('三個分點都有其他客戶，數字不全是焦家')
      and notes[-1].startswith('資料：富邦 DJ 分點進出（指定日期頁）'), notes)
check("抓不到的頁 (館前 10/02 存檔不完整 + 頁抓不到) → 「資料暫缺」註腳列出分點與日期",
      any(n.startswith('資料暫缺：元大-館前 10/02（') for n in notes), notes)

# ── D ──
print("\nD. 快取: 第二次產表不重抓")
cache = json.loads((DIR / tw.CACHE_NAME).read_text(encoding='utf-8'))
check("快取只存抓到的分點日 (失敗不存)", set((b, d) for b in cache['pages'] for d in cache['pages'][b])
      == set(PAGES), sorted((b, d) for b in cache['pages'] for d in cache['pages'][b]))
check("快取 585b 10/05 2492 = 金額 + 張數", cache['pages']['585b'][TD]['rows']['2492']['lots'] == [13, 20])
n1 = len(calls)
fetched_ok = {k for k in calls if k in PAGES}
wb2 = openpyxl.Workbook()
er._build_transfer_sheets(wb2, today, TD, DIR, fetch=True)
again = calls[n1:]
check("第二次 (可抓): 已快取的分點日一個都不重抓", not (set(again) & fetched_ok), again)
n2 = len(calls)
wb3 = openpyxl.Workbook()
er._build_transfer_sheets(wb3, today, TD, DIR, fetch=False)
check("最終重建 (fetch=False): 0 次抓取", len(calls) == n2, calls[n2:])
same = all(wb2['🔁 焦家匯撥'].cell(rr, cc).value == wb3['🔁 焦家匯撥'].cell(rr, cc).value
           for rr in range(1, 40) for cc in range(2, 10))
check("兩次產出的格子相同", same)
old_cap = tw.FETCH_CAP
tw.FETCH_CAP = 2
calls.clear()
er._build_transfer_sheets(openpyxl.Workbook(), today, TD, make_dir(), fetch=True)
tw.FETCH_CAP = old_cap
check("每次產表最多抓 FETCH_CAP 個分點日", len(calls) == 2, calls)

# ── E ──
print("\nE. 已存日檔必經 quarantine.filter_day (真的解密函式)")
qdir = pathlib.Path(tempfile.mkdtemp())
plain = json.dumps({'trade_date': '20260623', 'branches': [
    {'code': '9A9g', 'master': '陳族元', 'buys': [{'code': '2330'}], 'sells': [], 'error': None},
    {'code': '9216', 'master': '林滄海', 'buys': [{'code': '2603'}], 'sells': [], 'error': None}]})
pw = 'test-password'
(qdir / '20260623.json').write_text(
    json.dumps({'iterations': 1000, 'data': encrypt_data(plain, pw, iterations=1000)}), encoding='utf-8')
day = ORIG_READ(qdir, '20260623', pw)
by = {b['code']: b for b in day['branches']}
check("9A9g@0623 (data/quarantine.json) 被清空並標 quarantined",
      by['9A9g'].get('quarantined') is True and by['9A9g']['buys'] == [], by['9A9g'])
check("其他分點原樣", by['9216']['buys'] == [{'code': '2603'}])
check("沒有密碼 / 沒有檔案 → None", ORIG_READ(qdir, '20260623', '') is None and ORIG_READ(qdir, '20260624', pw) is None)
check("(C 已驗) 隔離中 / 失敗的分點日改查頁, 不當成「沒有進出」",
      ('585b', '20261001') in fetched_ok and ('984K', '20261001') in fetched_ok)

# ── F ──
print("\nF. 頁籤順序 (月檔)")
mdir = make_dir() / 'reports'
mdir.mkdir()
monthly = mdir / 'chip_radar_2026-10.xlsx'
wb0 = openpyxl.Workbook()
wb0.active.title = '20261002'
wb0.create_sheet(tw.SHEET_PREFIX + '已移除的群組')          # an unconfigured transfer sheet
wb0.save(monthly)
orig = (er.build_day_sheet, er.build_dashboard_sheet, er.build_pinned_track_sheet, er.build_mobile_summary_sheet)
er.build_day_sheet = lambda ws, *a, **k: 0
er.build_dashboard_sheet = lambda ws, *a, **k: None
er.build_pinned_track_sheet = lambda ws, *a, **k: None
er.build_mobile_summary_sheet = lambda ws, *a, **k: None
try:
    er._update_monthly_workbook(monthly, today, TD)
finally:
    er.build_day_sheet, er.build_dashboard_sheet, er.build_pinned_track_sheet, er.build_mobile_summary_sheet = orig
names = openpyxl.load_workbook(monthly).sheetnames
check("Dashboard → Pinned → 🔁 焦家匯撥 → 20261005 → 20261002; 舊群組頁被移除",
      names == [er.DASHBOARD_SHEET_NAME, er.PINNED_TRACK_SHEET_NAME, '🔁 焦家匯撥', '20261005', '20261002'], names)

# ── G ──
print("\nG. 真的抓頁路徑 (source_audit) + 假 http_get")


def page_html(echo, date, buy):
    head = ('<html><head><title>zgb0</title></head><body>\n<!-- ' + 'x' * 6000 + ' -->\n'
            f"<SCRIPT LANGUAGE=javascript><!--\n\tGenFundCorpCombo('{echo}','{echo}','frm');\n//--></SCRIPT>\n"
            f'<div class="t11">單位：仟元 ／ 資料日期：{date}</div>\n')
    rows = ''.join(f'<tr>\n<td class="t4t1" nowrap id="oAddCheckbox">\n'
                   f"<a href=\"javascript:Link2Stk('{c}');\">{c}{c}</a>\n</td>\n"
                   f'<td class="t3n1" nowrap>{v1:,}</td>\n<td class="t3n1" nowrap>{v2:,}</td>\n'
                   f'<td class="t3n1" nowrap>{v1 - v2:,}</td>\n</tr>\n' for c, v1, v2 in buy)
    return head + '<tr><td class="t2">買超</td></tr>\n' + rows + '<tr><td class="t2">賣超</td></tr>\n</table></body>\n</html>\n'


seen_urls = []


def fake_get(url, headers, timeout):
    seen_urls.append(url)
    echo = '0035003800350042' if 'wrongid' in TAG else '0035003800350062'
    lots = 'c=E' in url
    return 200, page_html(echo, TD, [('2492', 40 if lots else 12000, 10 if lots else 3000)]).encode('big5')


tw.fetch_branch_day = ORIG_FETCH
orig_get, orig_sleep = sa.http_get, sa._sleep
sa.http_get, sa._sleep = fake_get, (lambda s: None)
try:
    TAG = 'ok'
    d1 = make_dir()
    data = tw.collect([G], [], TD, d1, fetch=True, password='')
    c = data['焦家']['cells'][('585b', TD, '2492')]
    check("585b 以 hex 0035003800350062 送出, 指定日期 2026-10-5",
          any('a=0035003800350062' in u and 'e=2026-10-5&f=2026-10-5' in u for u in seen_urls), seen_urls[:2])
    check("頁面身分相符 → 40-10 = +30", c == {'state': 'value', 'lots': 30, 'est': False, 'src': 'page'}, c)
    TAG = 'wrongid'
    d2 = make_dir()
    data = tw.collect([G], [], TD, d2, fetch=True, password='')
    c = data['焦家']['cells'][('585b', TD, '2492')]
    check("頁面回 585B (統一-永和) 的身分 → 不收, 格子 missing", c['state'] == 'missing'
          and 'identity mismatch' in c.get('reason', ''), c)
    check("失敗不寫快取", not (d2 / tw.CACHE_NAME).exists())
finally:
    sa.http_get, sa._sleep = orig_get, orig_sleep
    tw.fetch_branch_day, tw.read_day = ORIG_FETCH, ORIG_READ

print()
print("─" * 72)
print(f"  整體: {'✅ ALL PASS' if all_pass else '❌ HAS FAIL'}")
sys.exit(0 if all_pass else 1)
