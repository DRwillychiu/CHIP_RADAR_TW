"""v3.80.22 mail rounds - which nightly run mails, and what its first line says.

Owner decision 2026-10-07: the two punctual runs (21:17 / 22:37, local Task
Scheduler -> workflow_dispatch) always mail; the late GitHub schedule re-runs
(02:00-06:00) mail only when the newest day sheet changed. The subject and
the first body line carry the round number and the change vs the previous
round, so an incomplete first mail can never pass as the final one.

State file data/reports/mail_rounds.json (committed with the data):
  {trade_date, rounds: [{round, at, event, rows, hash, mailed}],
   branches: {"<master>|<bno>": {name, rows, hash}}}   # last round only
Rows / content = the newest 8-digit day sheet of latest.xlsx, parsed exactly
like the source audit (src/audit/source_audit.parse_day_sheet), so the row
count matches the audit line of the same mail.

Writes send / round / line / tag to $GITHUB_OUTPUT. Always exits 0; on any
error send=true (a duplicate mail beats a missing one).

Usage:
  python scripts/mail_round.py --event schedule          # production
  python scripts/mail_round.py --event push --test       # TEST run: no state write
"""
import argparse
import hashlib
import json
import os
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

DEFAULT_XLSX = ROOT / "data" / "reports" / "latest.xlsx"
DEFAULT_STATE = ROOT / "data" / "reports" / "mail_rounds.json"
MAX_NAMES = 3


def _hash(items):
    raw = json.dumps(items, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]


def summarize(rows):
    """rows: parse_day_sheet() output -> (count, hash, {key: {name, rows, hash}})."""
    flat, by_branch = [], {}
    for x in rows:
        t = [x["master"], x["bno"], x["stock"], x["buy_lot"], x["sell_lot"],
             x["buy_wan"], x["sell_wan"]]
        flat.append(t)
        b = by_branch.setdefault(f"{x['master']}|{x['bno']}", {"name": x["branch"], "items": []})
        b["items"].append(t)
    branches = {k: {"name": v["name"], "rows": len(v["items"]), "hash": _hash(v["items"])}
                for k, v in by_branch.items()}
    return len(flat), _hash(flat), branches


def _changed_names(old, new):
    keys = list(new) + [k for k in old if k not in new]   # sheet order, then removed
    names = []
    for k in keys:
        if (old.get(k) or {}).get("hash") != (new.get(k) or {}).get("hash"):
            names.append((new.get(k) or old.get(k))["name"] or k)
    if len(names) > MAX_NAMES:
        return "、".join(names[:MAX_NAMES]) + f" 等 {len(names)} 個分點"
    return "、".join(names)


def decide(state, trade_date, count, digest, branches, event, test, now):
    """Pure decision. Returns (new_state, out) with out = {send, round, line, tag}."""
    same_day = isinstance(state, dict) and state.get("trade_date") == trade_date
    rounds = list(state.get("rounds") or []) if same_day else []
    old_branches = (state.get("branches") or {}) if same_day else {}
    prev = rounds[-1] if rounds else None
    n = len(rounds) + 1
    changed = prev is None or prev.get("hash") != digest
    send = bool(test) or event != "schedule" or changed

    if prev is None:
        desc, tag = "", f" · 第{n}輪"
    elif not changed:
        desc, tag = "，與上一輪相同", f" · 第{n}輪 無變化"
    else:
        diff = count - int(prev.get("rows") or 0)
        names = _changed_names(old_branches, branches)
        where = f"（變動：{names}）" if names else ""
        if diff:
            desc = f"，比上一輪{'多' if diff > 0 else '少'} {abs(diff)} 列{where}"
            tag = f" · 第{n}輪 {diff:+d}列"
        else:
            desc, tag = f"，列數相同但內容有變動{where}", f" · 第{n}輪 內容有變"
    line = f"🔁 第 {n} 輪（{now:%H:%M}）：{count} 列{desc}"

    rounds.append({"round": n, "at": now.isoformat(timespec="seconds"), "event": event,
                   "rows": count, "hash": digest, "mailed": send})
    new_state = {"trade_date": trade_date, "rounds": rounds, "branches": branches}
    return new_state, {"send": send, "round": n, "line": line, "tag": tag}


def _load_day(xlsx):
    from openpyxl import load_workbook
    from src.audit.source_audit import newest_day_sheet, parse_day_sheet
    wb = load_workbook(str(xlsx), read_only=True, data_only=True)
    try:
        name = newest_day_sheet(wb.sheetnames)
        if not name:
            raise ValueError("no 8-digit day sheet")
        return name, parse_day_sheet(wb[name])
    finally:
        wb.close()


def _write_outputs(out):
    path = os.environ.get("GITHUB_OUTPUT")
    if not path:
        return
    with open(path, "a", encoding="utf-8") as f:
        f.write(f"send={'true' if out['send'] else 'false'}\n")
        f.write(f"round={out['round']}\n")
        f.write(f"line={out['line']}\n")
        f.write(f"tag={out['tag']}\n")


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--event", default="workflow_dispatch")
    ap.add_argument("--test", action="store_true")
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    ap.add_argument("--state", default=str(DEFAULT_STATE))
    args = ap.parse_args(argv)
    state_path = Path(args.state)
    try:
        from src.audit.source_audit import TW
        trade_date, rows = _load_day(args.xlsx)
        count, digest, branches = summarize(rows)
        try:
            state = json.loads(state_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            state = {}
        new_state, out = decide(state, trade_date, count, digest, branches,
                                args.event, args.test, datetime.now(TW))
        if not args.test:
            state_path.parent.mkdir(parents=True, exist_ok=True)
            state_path.write_text(json.dumps(new_state, ensure_ascii=False, indent=1) + "\n",
                                  encoding="utf-8")
    except Exception as e:  # never block the mail
        out = {"send": True, "round": 0, "line": f"⚠️ 輪次判斷失敗（{type(e).__name__}），照常寄信",
               "tag": ""}
    print(out["line"])
    print(f"send={out['send']}")
    if not out["send"]:
        print("::notice title=Mail skipped::Schedule re-run, day sheet unchanged since the last round")
    _write_outputs(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
