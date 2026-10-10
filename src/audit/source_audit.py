"""
source_audit.py - v3.80.15 nightly source audit.

Every data row of the newest Excel day sheet is re-checked against the original
sources: the Fubon branch pages for that exact trade date (zgb0, date pinned,
full page) and the official closes (TWSE MI_INDEX + TPEx daily quotes).

Why: auto_audit.py only does sanity checks on the Excel itself. The one-off
cross-check of 2026-09-25 caught bugs it cannot see (one branch served another
branch's page; pages cut at 30 rows so real lots were replaced by estimates).

Every verdict rule mirrors how the crawler fills the cell (keep in sync):
  src/pipelines/crawler_fetch.py fetch_branch_combined / merge_rows:
      a row of the net-buy list takes amounts (thousand TWD) from the amount page
      (c=B) and lots from the lots page (c=E), both looked up by stock code in
      the SAME region (net-buy) of the full page. amt_listed / lot_listed = the
      stock is in that region of that page.
  crawler.py estimate block (v3.27.2 / v3.80.1): only a field whose page does
      not list the stock is estimated from the close, and only if close > 0:
          lots  = max(1, round(amt_k / close))   when amt_k > 0, else 0
          amt_k = int(round(lots * close))       when lots > 0, else 0
  src/exports/excel_report.py _write_stock_row:
      E/F = lots as stored, G/H (wan) = round(amt_k / 10) if amt_k else 0.
  excel_report._top_stocks_for_branch: buys first, then sells, deduped by code,
      net > 0 only -> every sheet row comes from the net-buy list; the net-sell
      region is only used as a fallback lookup.

Row verdicts:
  REAL        stock on both pages, lots and amounts equal the page values
  LOT_EST     amounts equal the amount page, stock not on the lots page, lots
              equal the crawler's estimate from amount / official close
  AMT_EST     lots equal the lots page, stock not on the amount page, amounts
              equal the crawler's estimate from lots x official close
  MISMATCH    anything else
  UNVERIFIED  the branch's source pages could not be fetched / validated
              (identity, data date, truncated page, time budget), or an
              estimate row has no official close to check against

Pure functions (sheet parser, page parser, row classifier, email line) have no
I/O. The fetch layer goes through http_get() only, so tests replace that one
function and run fully offline.
"""
import re
import time
from datetime import datetime, timedelta, timezone

from src.pipelines.crawler_fetch import (
    HOME_URL, IDENTITY_RE, UA_POOL, fubon_bno, parse_region,
)

FUBON_URL = ("https://fubon-ebrokerdj.fbs.com.tw/z/zg/zgb/zgb0.djhtm"
             "?a={b}&b={b}&c={mode}&e={d}&f={d}")
TWSE_URL = ("https://www.twse.com.tw/rwd/zh/afterTrading/MI_INDEX"
            "?date={ymd}&type=ALLBUT0999&response=json")
TPEX_URL = ("https://www.tpex.org.tw/www/zh-tw/afterTrading/otc"
            "?date={y}/{m}/{d}&type=EW&response=json")
# v3.80.18: fallback for the TPEx website, which did not answer the GitHub
# runner on 2026-10-06 (34 rows left UNVERIFIED). Same endpoint the crawler
# uses (fetchers/institutional.py TPEX_DAILY_URL); latest day only, so the
# Date field is checked. ~2 MB, slow -> longer timeout.
TPEX_OPENAPI_URL = "https://www.tpex.org.tw/openapi/v1/tpex_mainboard_daily_close_quotes"
TPEX_OPENAPI_TIMEOUT_S = 60

_UA = UA_POOL[0]   # full Chrome UA, same as the crawler
FUBON_HEADERS = {
    "User-Agent": _UA,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "zh-TW,zh;q=0.9,en;q=0.8",
    "Referer": HOME_URL,
}
TWSE_HEADERS = {"User-Agent": _UA, "Accept": "application/json, text/plain, */*",
                "Referer": "https://www.twse.com.tw/"}
TPEX_HEADERS = {"User-Agent": _UA, "Accept": "application/json, text/plain, */*",
                "Referer": "https://www.tpex.org.tw/"}

# v3.85.6 (owner 2026-10-09): 1.1 -> the shared 1.6 s floor (light_verify and
# transfer_watch use this pacer too)
from src.core.fubon_pacing import FUBON_MIN_GAP_S, note_request  # noqa: E402
from src.core.fubon_codec import decode_fubon  # noqa: E402
FUBON_GAP_S = FUBON_MIN_GAP_S   # pause between two Fubon requests
FUBON_TIMEOUT_S = 20       # per request (crawler uses 20 too)
CLOSE_TIMEOUT_S = 20
PAGE_ATTEMPTS = 2          # one retry per page while the budget lasts
BUDGET_S = 480             # whole audit; later branches become UNVERIFIED
MAX_MISMATCHES = 50        # listed in the json (the count is always complete)

VERDICTS = ("REAL", "LOT_EST", "AMT_EST", "MISMATCH", "UNVERIFIED")
DATA_DATE_RE = re.compile(r"資料日期：(\d{8})")
_STOCK_RE = re.compile(r"^(?P<name>.*)\((?P<code>[0-9A-Za-z]+)\)\s*$")
_HEADER_D = "標的"
TW = timezone(timedelta(hours=8))


class SourceError(Exception):
    """A source page / close list could not be fetched or failed validation."""


class BudgetExhausted(SourceError):
    pass


# ----------------------------------------------------------------------------
#  Formulas mirrored from the crawler / Excel builder (see module docstring)
# ----------------------------------------------------------------------------

def est_lots(amt_k, close):
    """crawler.py: s["buy_lot"] = max(1, round(buy_amt_k / cp_close)) if amt > 0."""
    return max(1, round(amt_k / close)) if amt_k > 0 else 0


def est_amt_k(lots, close):
    """crawler.py: s["buy_amt"] = int(round(buy_lot_raw * cp_close)) if lots > 0."""
    return int(round(lots * close)) if lots > 0 else 0


def to_wan(amt_k):
    """excel_report._write_stock_row: round(buy_amt_k / 10) if buy_amt_k else 0."""
    return round(amt_k / 10) if amt_k else 0


# ----------------------------------------------------------------------------
#  Day sheet parser
# ----------------------------------------------------------------------------

def newest_day_sheet(sheetnames):
    days = sorted(n for n in sheetnames if len(n) == 8 and n.isdigit())
    return days[-1] if days else None


def _cell_int(v):
    if v is None or v == "":
        return 0
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else None
    try:
        f = float(str(v).replace(",", "").strip())
    except ValueError:
        return None
    return int(f) if f.is_integer() else None


def parse_day_rows(rows):
    """rows: iterable of row tuples (columns A..H at least), sheet order.

    Layout written by excel_report.build_day_sheet: A master (merged over its
    block, value on the first data row), B branch name, C branch code (merged
    over the branch rows), D 'name(code)', E buy lots, F sell lots, G buy wan,
    H sell wan. Header rows carry the column label in D (_HEADER_D); notice
    rows and padding rows have no '(code)' in D; the optional top warning row
    starts with a warning sign; the footer row starts with the info sign
    (U+24D8) and ends the blocks.
    """
    out = []
    master = branch = bno = None
    for r, vals in enumerate(rows, start=1):
        vals = (list(vals) + [None] * 8)[:8]
        a, b, c, d, e, f, g, h = vals
        if isinstance(a, str) and a.startswith("ⓘ"):
            break
        if isinstance(d, str) and d.strip() == _HEADER_D:
            continue
        if isinstance(a, str) and a.strip() and not a.startswith("⚠"):
            master = a.strip()
        if b not in (None, ""):
            branch = str(b).strip()
        if c not in (None, ""):
            bno = str(c).strip()
        m = _STOCK_RE.match(d.strip()) if isinstance(d, str) else None
        if not m or not bno:
            continue
        out.append({
            "row": r, "master": master, "branch": branch, "bno": bno,
            "stock": d.strip(), "name": m.group("name"), "code": m.group("code"),
            "buy_lot": _cell_int(e), "sell_lot": _cell_int(f),
            "buy_wan": _cell_int(g), "sell_wan": _cell_int(h),
        })
    return out


def parse_day_sheet(ws):
    """ws: openpyxl worksheet (normal or read-only)."""
    return parse_day_rows(ws.iter_rows(min_col=1, max_col=8, values_only=True))


# ----------------------------------------------------------------------------
#  Source page parsers
# ----------------------------------------------------------------------------

def fubon_date(trade_date):
    """'20261005' -> '2026-10-5' (the form Fubon's own links use)."""
    return f"{int(trade_date[:4])}-{int(trade_date[4:6])}-{int(trade_date[6:8])}"


def fubon_url(branch_code, mode, trade_date):
    b = fubon_bno(branch_code)
    return FUBON_URL.format(b=b, mode=mode, d=fubon_date(trade_date))


def parse_fubon_page(html, sent, trade_date):
    """-> {'buy': [rows], 'sell': [rows]} (full page, parse_region rows).

    Rejects what the crawler would reject (too small, other branch's identity)
    plus what only a date-pinned page can prove: the data date, and a page
    that was cut off before </html>. A page without the id marker cannot be
    validated, so it is rejected too (the crawler only warns)."""
    if len(html) < 5000:
        raise SourceError(f"page too small ({len(html)} chars)")
    if "</html>" not in html[-3000:].lower():
        raise SourceError("page truncated (no </html>)")
    idm = IDENTITY_RE.search(html)
    if not idm:
        raise SourceError("page has no branch-id marker")
    if idm.group(2) != sent:
        raise SourceError(f"identity mismatch: sent {sent}, page is {idm.group(2)}")
    dm = DATA_DATE_RE.search(html)
    if not dm:
        raise SourceError("page has no data date")
    if dm.group(1) != trade_date:
        raise SourceError(f"date mismatch: asked {trade_date}, page is {dm.group(1)}")
    i, j = html.find("買超</td>"), html.find("賣超</td>")
    if i < 0 or j < i:
        raise SourceError("net-buy / net-sell tables not found")
    return {"buy": parse_region(html[i:j]), "sell": parse_region(html[j:])}


def page_maps(page):
    """Same lookup the crawler builds: {code: (v1, v2)} per region, full page."""
    return {reg: {r["code"]: (r["v1"], r["v2"]) for r in page[reg]}
            for reg in ("buy", "sell")}


def _price(s):
    try:
        v = float(str(s).replace(",", "").strip())
    except ValueError:
        return None
    return v if v > 0 else None


def parse_twse_closes(doc, trade_date):
    if not isinstance(doc, dict) or doc.get("stat") != "OK":
        raise SourceError(f"TWSE stat {doc.get('stat') if isinstance(doc, dict) else type(doc).__name__}")
    if str(doc.get("date")) != trade_date:
        raise SourceError(f"TWSE date mismatch: asked {trade_date}, got {doc.get('date')}")
    out = {}
    for tb in doc.get("tables") or []:
        f = [str(x).strip() for x in (tb.get("fields") or [])]
        if "證券代號" in f and "收盤價" in f:
            ci, pi = f.index("證券代號"), f.index("收盤價")
            for r in tb.get("data") or []:
                c = _price(r[pi]) if len(r) > pi else None
                if c:
                    out[str(r[ci]).strip()] = c
    if not out:
        raise SourceError("TWSE close table not found")
    return out


def parse_tpex_closes(doc, trade_date):
    if not isinstance(doc, dict):
        raise SourceError(f"TPEx bad payload {type(doc).__name__}")
    if str(doc.get("date")) != trade_date:
        raise SourceError(f"TPEx date mismatch: asked {trade_date}, got {doc.get('date')}")
    out = {}
    for tb in doc.get("tables") or []:
        f = [str(x).strip() for x in (tb.get("fields") or [])]
        if "代號" in f and "收盤" in f:
            ci, pi = f.index("代號"), f.index("收盤")
            for r in tb.get("data") or []:
                c = _price(r[pi]) if len(r) > pi else None
                if c:
                    out[str(r[ci]).strip()] = c
    if not out:
        raise SourceError("TPEx close table not found")
    return out


def _roc_date(trade_date):
    """'20261005' -> '1151005' (TPEx OpenAPI Date field)."""
    return f"{int(trade_date[:4]) - 1911}{trade_date[4:]}"


def parse_tpex_openapi_closes(doc, trade_date):
    if not isinstance(doc, list) or not doc or not isinstance(doc[0], dict):
        raise SourceError(f"TPEx OpenAPI bad payload {type(doc).__name__}")
    want, got = _roc_date(trade_date), str(doc[0].get("Date") or "").strip()
    if got != want:
        raise SourceError(f"TPEx OpenAPI date mismatch: asked {want}, got {got}")
    out = {}
    for item in doc:
        code = str(item.get("SecuritiesCompanyCode") or "").strip()
        c = _price(item.get("Close"))
        if code and c:
            out[code] = c
    if not out:
        raise SourceError("TPEx OpenAPI has no closes")
    return out


# ----------------------------------------------------------------------------
#  Row classifier
# ----------------------------------------------------------------------------

def lookup(maps_b, maps_e, code):
    """-> (region, amt (buy, sell) or None, lots (buy, sell) or None).
    Net-buy region first (that is where every sheet row comes from)."""
    for reg in ("buy", "sell"):
        a, l = maps_b[reg].get(code), maps_e[reg].get(code)
        if a is not None or l is not None:
            return reg, a, l
    return None, None, None


_FIELDS = ("buy_lot", "sell_lot", "buy_wan", "sell_wan")
# user-facing labels (json 'reason' -> GitHub Issue)
_LABEL = {"buy_lot": "買張", "sell_lot": "賣張", "buy_wan": "買萬", "sell_wan": "賣萬"}
_KIND = {"REAL": "真實值", "LOT_EST": "估算張數", "AMT_EST": "估算金額"}
R_NOT_ON_PAGE = "富邦原始頁 (買超/賣超) 都沒有這檔"
R_NO_CLOSE = "沒有官方收盤價, 無法驗證估算值"
H_LOTS_EST = " (Excel 張數 = 收盤價估算值: 抓取當時張數頁沒有這檔)"
H_AMT_EST = " (Excel 金額 = 收盤價估算值: 抓取當時金額頁沒有這檔)"
H_ZERO = " (Excel 為 0: 爬蟲當時沒有這檔的收盤價)"


def classify_row(x, amt, lots, close):
    """x: parsed sheet row. amt / lots: (buy, sell) from the same region of the
    amount / lots page, or None when the stock is not there.
    -> (verdict, expected dict or None, reason or None)."""
    if amt is None and lots is None:
        return "MISMATCH", None, R_NOT_ON_PAGE
    if amt is not None and lots is not None:
        kind = "REAL"
        exp = (lots[0], lots[1], to_wan(amt[0]), to_wan(amt[1]))
    elif not close:
        return "UNVERIFIED", None, R_NO_CLOSE
    elif amt is not None:
        kind = "LOT_EST"
        exp = (est_lots(amt[0], close), est_lots(amt[1], close), to_wan(amt[0]), to_wan(amt[1]))
    else:
        kind = "AMT_EST"
        exp = (lots[0], lots[1], to_wan(est_amt_k(lots[0], close)), to_wan(est_amt_k(lots[1], close)))
    expected = dict(zip(_FIELDS, exp))
    got = tuple(x.get(k) for k in _FIELDS)
    if got == exp:
        return kind, expected, None
    diff = [k for k, g, e in zip(_FIELDS, got, exp) if g != e]
    reason = f"{_KIND[kind]}規則: " + ", ".join(f"{_LABEL[k]} {x.get(k)} != {expected[k]}" for k in diff)
    # hints for the reader: what the differing sheet values look like instead
    if kind == "REAL" and close:
        est = {"buy_lot": est_lots(amt[0], close), "sell_lot": est_lots(amt[1], close),
               "buy_wan": to_wan(est_amt_k(lots[0], close)), "sell_wan": to_wan(est_amt_k(lots[1], close))}
        same = all(x.get(k) == est[k] for k in diff)
        if same and set(diff) <= {"buy_lot", "sell_lot"}:
            reason += H_LOTS_EST
        elif same and set(diff) <= {"buy_wan", "sell_wan"}:
            reason += H_AMT_EST
    est_fields = {"LOT_EST": {"buy_lot", "sell_lot"}, "AMT_EST": {"buy_wan", "sell_wan"}}.get(kind)
    if est_fields and set(diff) <= est_fields and all(x.get(k) == 0 for k in diff):
        reason += H_ZERO
    return "MISMATCH", expected, reason


# ----------------------------------------------------------------------------
#  Fetch layer (network only through http_get)
# ----------------------------------------------------------------------------

_session = None
# time hooks, resolved at call time (tests replace them to run instantly)
_sleep = time.sleep
_clock = time.monotonic


def http_get(url, headers, timeout):
    """The only network call. -> (status_code, body bytes)."""
    global _session
    import requests
    if _session is None:
        _session = requests.Session()
    r = _session.get(url, headers=headers, timeout=timeout)
    return r.status_code, r.content


def _get_json(get, url, headers, attempts=2, timeout=CLOSE_TIMEOUT_S):
    import json
    last = None
    for _ in range(attempts):
        try:
            status, body = get(url, headers, timeout)
            if status != 200:
                last = f"HTTP {status}"
                continue
            return json.loads(body.decode("utf-8-sig", errors="replace") if isinstance(body, bytes) else body)
        except Exception as e:      # network / JSON error -> retry once
            last = f"{type(e).__name__}: {e}"
    raise SourceError(last or "no response")


def fetch_closes(trade_date, get=None):
    """-> (closes {code: close}, errors [str]). TWSE overrides TPEx like the crawler."""
    get = get or http_get
    y, m, d = trade_date[:4], trade_date[4:6], trade_date[6:8]
    twse, tpex, errors = {}, {}, []
    try:
        twse = parse_twse_closes(_get_json(get, TWSE_URL.format(ymd=trade_date), TWSE_HEADERS), trade_date)
    except SourceError as e:
        errors.append(f"TWSE: {e}")
    try:
        tpex = parse_tpex_closes(_get_json(get, TPEX_URL.format(y=y, m=m, d=d), TPEX_HEADERS), trade_date)
    except SourceError as e:
        try:
            tpex = parse_tpex_openapi_closes(
                _get_json(get, TPEX_OPENAPI_URL, TPEX_HEADERS, timeout=TPEX_OPENAPI_TIMEOUT_S), trade_date)
            print(f"  TPEx website failed ({e}); OpenAPI fallback: {len(tpex)} closes")
        except SourceError as e2:
            errors.append(f"TPEx: {e}; OpenAPI fallback: {e2}")
    return {**tpex, **twse}, errors


class _Pacer:
    """Fubon pacing + whole-audit time budget."""

    def __init__(self, sleep, clock, deadline, gap_s=FUBON_GAP_S):
        self.sleep, self.clock, self.deadline, self.gap_s = sleep, clock, deadline, gap_s
        self.n = 0

    def before_request(self):
        if self.clock() >= self.deadline:
            raise BudgetExhausted("time budget exhausted")
        if self.n:
            self.sleep(self.gap_s)
        self.n += 1
        note_request()                  # v3.85.6: measured gap


def fetch_branch_pages(branch_code, trade_date, get, pacer):
    """-> {'B': page_maps, 'E': page_maps}; raises SourceError."""
    sent = fubon_bno(branch_code)
    pages = {}
    for mode in ("B", "E"):
        url = fubon_url(branch_code, mode, trade_date)
        last = None
        for _ in range(PAGE_ATTEMPTS):
            pacer.before_request()
            try:
                status, body = get(url, FUBON_HEADERS, FUBON_TIMEOUT_S)
            except Exception as e:
                last = f"{type(e).__name__}: {e}"
                continue
            if status != 200:
                last = f"HTTP {status}"
                continue
            html = decode_fubon(body) if isinstance(body, bytes) else body   # v3.85.8 cp950
            try:
                pages[mode] = page_maps(parse_fubon_page(html, sent, trade_date))
                break
            except SourceError as e:
                last = str(e)
        if mode not in pages:
            raise SourceError(f"{'amount' if mode == 'B' else 'lots'} page: {last}")
    return pages


# ----------------------------------------------------------------------------
#  Audit
# ----------------------------------------------------------------------------

def _now_tw():
    return datetime.now(TW).isoformat(timespec="seconds")


def audit_rows(rows, trade_date, closes, get=None, sleep=None, clock=None,
               deadline=None, log=print):
    """Fetch every branch of `rows` and classify each row.
    -> (verdict list aligned with rows, unverified_branches)."""
    get = get or http_get
    sleep = sleep or _sleep
    clock = clock or _clock
    pacer = _Pacer(sleep, clock, deadline if deadline is not None else clock() + BUDGET_S)
    order = list(dict.fromkeys(x["bno"] for x in rows))
    pages, failed = {}, {}
    for i, bno in enumerate(order, 1):
        if failed.get("_budget"):
            failed[bno] = failed["_budget"]
            continue
        try:
            pages[bno] = fetch_branch_pages(bno, trade_date, get, pacer)
            log(f"  [{i}/{len(order)}] {bno} ok")
        except BudgetExhausted as e:
            failed[bno] = failed["_budget"] = str(e)
            log(f"  [{i}/{len(order)}] {bno} UNVERIFIED: {e}")
        except SourceError as e:
            failed[bno] = str(e)
            log(f"  [{i}/{len(order)}] {bno} UNVERIFIED: {e}")
    failed.pop("_budget", None)

    results = []
    for x in rows:
        bno = x["bno"]
        if bno in failed:
            results.append({"verdict": "UNVERIFIED", "reason": f"source page: {failed[bno]}",
                            "unverified": "page"})
            continue
        reg, amt, lots = lookup(pages[bno]["B"], pages[bno]["E"], x["code"])
        close = closes.get(x["code"])
        verdict, expected, reason = classify_row(x, amt, lots, close)
        res = {"verdict": verdict, "region": reg, "amt_k": list(amt) if amt else None,
               "lots": list(lots) if lots else None, "close": close,
               "expected": expected, "reason": reason}
        if verdict == "UNVERIFIED":
            res["unverified"] = "close"
        results.append(res)

    unverified_branches = []
    for bno in order:
        if bno in failed:
            mine = [x for x in rows if x["bno"] == bno]
            unverified_branches.append({
                "bno": bno, "master": mine[0]["master"], "branch": mine[0]["branch"],
                "rows": len(mine), "reason": failed[bno]})
    return results, unverified_branches


def run_audit(xlsx_path, sheet=None, get=None, sleep=None, clock=None,
              budget_s=BUDGET_S, log=print):
    """Audit one day sheet (default: the newest). Never raises: problems end
    up in status 'error' with the message in 'error'."""
    clock = clock or _clock
    t0 = clock()
    doc = {"trade_date": None, "sheet": None, "xlsx": str(xlsx_path),
           "checked_at": _now_tw(), "status": "error", "rows": 0, "branches": 0,
           "verdicts": {v: 0 for v in VERDICTS}, "mismatch_total": 0, "mismatches": [],
           "unverified_branches": [], "unverified_rows_page": 0,
           "unverified_rows_close": 0, "closes_loaded": 0, "close_errors": [],
           "elapsed_s": None, "error": None}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
        try:
            name = sheet or newest_day_sheet(wb.sheetnames)
            if not name or name not in wb.sheetnames:
                raise SourceError(f"day sheet not found ({sheet or 'no 8-digit sheet'})")
            rows = parse_day_sheet(wb[name])
        finally:
            wb.close()
        doc["sheet"] = doc["trade_date"] = name
        doc["rows"] = len(rows)
        doc["branches"] = len({x["bno"] for x in rows})
        if not rows:
            raise SourceError(f"no data rows in day sheet {name}")
        log(f"[source audit] {name}: {len(rows)} rows / {doc['branches']} branches")

        closes, close_errors = fetch_closes(name, get)
        doc["closes_loaded"], doc["close_errors"] = len(closes), close_errors
        log(f"  official closes: {len(closes)}" + (f" | errors: {close_errors}" if close_errors else ""))

        results, unv = audit_rows(rows, name, closes, get=get, sleep=sleep, clock=clock,
                                  deadline=t0 + budget_s, log=log)
        mism = []
        for x, res in zip(rows, results):
            doc["verdicts"][res["verdict"]] += 1
            if res.get("unverified") == "page":
                doc["unverified_rows_page"] += 1
            elif res.get("unverified") == "close":
                doc["unverified_rows_close"] += 1
            if res["verdict"] == "MISMATCH":
                mism.append({
                    "row": x["row"], "master": x["master"], "branch": x["branch"],
                    "bno": x["bno"], "stock": x["stock"], "code": x["code"],
                    "excel": {k: x[k] for k in _FIELDS},
                    "expected": res["expected"],
                    "source": {"region": res["region"], "amt_k": res["amt_k"],
                               "lots": res["lots"], "close": res["close"]},
                    "reason": res["reason"]})
        doc["mismatch_total"] = len(mism)
        doc["mismatches"] = mism[:MAX_MISMATCHES]
        doc["unverified_branches"] = unv
        v = doc["verdicts"]
        doc["status"] = "mismatch" if v["MISMATCH"] else ("incomplete" if v["UNVERIFIED"] else "ok")
    except Exception as e:
        doc["status"] = "error"
        doc["error"] = f"{type(e).__name__}: {e}"
    doc["elapsed_s"] = round(clock() - t0, 1)
    return doc


# ----------------------------------------------------------------------------
#  Email line (one Traditional Chinese line at the top of the daily mail)
# ----------------------------------------------------------------------------

def email_line(doc, latest):
    """doc: the audit json (or None); latest: data/index.json 'latest'."""
    if not isinstance(doc, dict) or not latest or doc.get("trade_date") != latest:
        return "⚠️ 來源比對：今天沒有跑完"
    td = str(doc["trade_date"])
    md = f"{td[4:6]}/{td[6:8]}"
    v = doc.get("verdicts") or {}
    st = doc.get("status")
    unv = v.get("UNVERIFIED", 0)
    if st == "ok":
        parts = [f"真實 {v.get('REAL', 0)}"]
        if v.get("LOT_EST"):
            parts.append(f"估算張數 {v['LOT_EST']}")
        if v.get("AMT_EST"):
            parts.append(f"估算金額 {v['AMT_EST']}")
        return f"✅ 來源比對 {md}：{doc.get('rows', 0)} 列全部相符（{'、'.join(parts)}）"
    if st == "mismatch":
        n = doc.get("mismatch_total") or v.get("MISMATCH", 0)
        tail = f"（另有 {unv} 列未比對）" if unv else ""
        return f"🚨 來源比對 {md}：{n} 列不符，詳見 GitHub Issue{tail}"
    if st == "incomplete":
        why = []
        nb = len(doc.get("unverified_branches") or [])
        if nb:
            why.append(f"{nb} 個分點抓不到原始頁")
        if doc.get("unverified_rows_close"):
            why.append(f"{doc['unverified_rows_close']} 列缺官方收盤價")
        return f"⚠️ 來源比對 {md}：{'、'.join(why) or '部分資料無法比對'}，未比對 {unv} 列"
    return f"⚠️ 來源比對 {md}：比對程式出錯，詳見 Actions 紀錄"
